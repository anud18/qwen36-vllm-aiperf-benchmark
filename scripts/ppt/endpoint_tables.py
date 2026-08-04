import json, os
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

INK=RGBColor(0x26,0x26,0x2a); MUTED=RGBColor(0x6b,0x6b,0x70); ENDPT=RGBColor(0xeb,0x68,0x34)
SURF=RGBColor(0xfc,0xfc,0xfb); HEADBG=RGBColor(0xf2,0xf2,0xf0); WARN=RGBColor(0xe3,0x49,0x48)
FONT="Noto Sans CJK TC"

# (model, point label, requested, artifact dir, note)
EP=[("Qwen3-VL-30B-A3B","chatbot_flat c1  (warmup 2)",10,"qwen3vl_ngrok/chatbot_flat/c1",None),
 ("Qwen3-VL-30B-A3B","chatbot_flat c2  (warmup 2)",20,"qwen3vl_ngrok/chatbot_flat/c2",None),
 ("Qwen3-VL-30B-A3B","chatbot_flat c1 重跑  (warmup 2)",10,"qwen3vl_ngrok_run2/chatbot_flat/c1",None),
 ("Qwen3-VL-30B-A3B","agent_flat c1  (warmup 2)",10,"qwen3vl_ngrok/agent_flat/c1",None),
 ("Qwen3-VL-30B-A3B","chatbot_flat c1  (no warmup)",20,"w0_chatbot/chatbot_flat/c1",None),
 ("Qwen3-VL-30B-A3B","agent_flat c1  (no warmup)",20,"w0_agent/agent_flat/c1",None),
 ("Qwen3-VL-30B-A3B","chatbot_flat c1 重跑  (no warmup)",20,"w0_chatbot_rep/chatbot_flat/c1",None),
 ("Qwen3-VL-30B-A3B","agent_flat c2  (no warmup)",20,"w0_agent/agent_flat/c2",
  "無效：20 筆中 6 筆撞上 ngrok 月流量配額（403 ERR_NGROK_725）。剩餘 14 筆的併發時序也被拖累，不可引用。"),
 ("Llama-3.1-8B","chatbot_flat c1  (no warmup)",20,"cf_llama31_chatbot/chatbot_flat/c1",None),
 ("Llama-3.1-8B","agent_flat c1  (no warmup)",20,"cf_llama31_agent/agent_flat/c1",None),
 ("Llama-3.1-8B","chatbot_flat c1 重跑  (worker 重啟後)",20,"cf_llama31_chatbot_rep/chatbot_flat/c1",None),
 ("Llama-3.1-8B","chatbot_flat c1  (100 請求)",100,"r100_chatbot_flat_c1/chatbot_flat/c1",None),
 ("Llama-3.1-8B","chatbot_flat c2  (100 請求)",100,"r100_chatbot_flat_c2/chatbot_flat/c2",None),
 ("Llama-3.1-8B","agent_flat c1  (100 請求)",100,"r100_agent_flat_c1/agent_flat/c1",
  "部分無效：100 筆中 3 筆 Cloudflare 524 逾時。524 砍掉的是最慢的請求，故下列數字偏樂觀。"),
 ("Llama-3.1-8B","agent_flat c2  (100 請求)",100,"r100_agent_flat_c2/agent_flat/c2",
  "部分無效：100 筆中 10 筆 Cloudflare 524 逾時。524 砍掉的是最慢的請求，故下列數字偏樂觀。")]

ROWS=[("TTFT (ms)","time_to_first_token",1),("ITL = TPOT (ms)","inter_token_latency",2),
      ("E2E latency (ms)","request_latency",1),
      ("Prefill throughput (tok/s/user)","prefill_throughput_per_user",1),
      ("Decode throughput (tok/s/user)","output_token_throughput_per_user",2),
      ("ISL (tokens)","input_sequence_length",1),("OSL (tokens)","output_sequence_length",1)]
g=lambda r,k,f='avg': (r[k].get(f) if isinstance(r[k],dict) else r[k]) if k in r else None

def add_endpoint_slides(prs, blank, head, tb, rule):
    s=blank(); head(s,"附錄：Dynamo 端點原始數據","逐點明細")
    tb(s,0.9,2.3,11.5,2.4,
     "以下 15 頁為端點側每個測試點的完整分布，含 avg / P50 / P90 / P95。\n\n"
     "3 個點標記為無效或部分無效 —— 皆為 tunnel 層問題（ngrok 流量配額、Cloudflare 逾時），\n"
     "非伺服器行為。逾時砍掉的是最慢的請求，會讓數字看起來比實際更好，故不納入任何結論。",16,False,INK,space=9)
    for mdl,name,req,d,note in EP:
        p=f"results/nvfp4/{d}/profile_export_aiperf.json"
        if not os.path.exists(p): continue
        R=json.load(open(p))
        s=blank(); head(s,f"{mdl} · {name}","附錄：Dynamo 端點")
        n=len(ROWS)+1
        t=s.shapes.add_table(n,5,Inches(0.9),Inches(2.05),Inches(11.5),Inches(0.42*n)).table
        t.columns[0].width=Inches(3.9)
        for c in range(1,5): t.columns[c].width=Inches(1.9)
        for c,txt in enumerate(["指標","avg","P50","P90","P95"]):
            cell=t.cell(0,c); cell.text=""
            pr=cell.text_frame.paragraphs[0]; r=pr.add_run(); r.text=txt
            r.font.size=Pt(13); r.font.bold=True; r.font.name=FONT
            r.font.color.rgb=ENDPT if c>0 else INK
            pr.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
            cell.fill.solid(); cell.fill.fore_color.rgb=HEADBG
        for i,(lbl,key,nd) in enumerate(ROWS,start=1):
            vals=[lbl]+[f"{g(R,key,f):,.{nd}f}" for f in ("avg","p50","p90","p95")]
            for c,txt in enumerate(vals):
                cell=t.cell(i,c); cell.text=""
                pr=cell.text_frame.paragraphs[0]; r=pr.add_run(); r.text=txt
                r.font.size=Pt(13); r.font.name=FONT; r.font.color.rgb=INK
                pr.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
                cell.fill.solid(); cell.fill.fore_color.rgb=SURF
        y=2.05+0.42*n+0.22
        ok=int(g(R,'request_count'))
        line=(f"Request number  {req}          Successful requests  {ok} / {req}          "
              f"Request throughput  {g(R,'request_throughput'):.4f} req/s          "
              f"ISL total  {g(R,'total_isl'):,.0f}          OSL total  {g(R,'total_osl'):,.0f}")
        tb(s,0.9,y,11.5,0.5,line,12,True,INK)
        cache=g(R,'overall_usage_prompt_cache_read_pct')
        tb(s,0.9,y+0.42,11.5,0.5,
           f"Prefix cache 命中率 {cache:.2f}%          OSL 偏差 {g(R,'osl_mismatch_count'):.0f}          "
           f"benchmark duration {g(R,'benchmark_duration'):,.1f} s",12,False,MUTED)
        if note: tb(s,0.9,y+0.92,11.5,0.7,"⚠  "+note,12,True,WARN,space=3)
