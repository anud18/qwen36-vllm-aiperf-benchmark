import json
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

INK=RGBColor(0x26,0x26,0x2a); MUTED=RGBColor(0x6b,0x6b,0x70); SPARK=RGBColor(0x2a,0x78,0xd6)
ENDPT=RGBColor(0xeb,0x68,0x34); SURF=RGBColor(0xfc,0xfc,0xfb); RULE=RGBColor(0xe4,0xe4,0xe2)
HEADBG=RGBColor(0xf2,0xf2,0xf0); FONT="Noto Sans CJK TC"

PAIRS=[("Qwen3-VL-30B-A3B","w2 chatbot c1",10,"qwen3vl","w2_chatbot_c1/chatbot_flat/c1","qwen3vl_ngrok/chatbot_flat/c1"),
 ("Qwen3-VL-30B-A3B","w2 chatbot c2",20,"qwen3vl","w2_chatbot_c2/chatbot_flat/c2","qwen3vl_ngrok/chatbot_flat/c2"),
 ("Qwen3-VL-30B-A3B","w2 chatbot c1 (repeat)",10,"qwen3vl","w2_chatbot_c1_rep/chatbot_flat/c1","qwen3vl_ngrok_run2/chatbot_flat/c1"),
 ("Qwen3-VL-30B-A3B","w2 agent c1",10,"qwen3vl","w2_agent_c1/agent_flat/c1","qwen3vl_ngrok/agent_flat/c1"),
 ("Qwen3-VL-30B-A3B","w0 chatbot c1",20,"qwen3vl","w0_chatbot_c1/chatbot_flat/c1","w0_chatbot/chatbot_flat/c1"),
 ("Qwen3-VL-30B-A3B","w0 agent c1",20,"qwen3vl","w0_agent_c1/agent_flat/c1","w0_agent/agent_flat/c1"),
 ("Qwen3-VL-30B-A3B","w0 chatbot c1 (repeat)",20,"qwen3vl","w0_chatbot_c1_rep/chatbot_flat/c1","w0_chatbot_rep/chatbot_flat/c1"),
 ("Llama-3.1-8B","w0 chatbot c1",20,"llama31","w0_chatbot_c1/chatbot_flat/c1","cf_llama31_chatbot/chatbot_flat/c1"),
 ("Llama-3.1-8B","w0 agent c1",20,"llama31","w0_agent_c1/agent_flat/c1","cf_llama31_agent/agent_flat/c1"),
 ("Llama-3.1-8B","w0 chatbot c1 (repeat)",20,"llama31","w0_chatbot_c1_rep/chatbot_flat/c1","cf_llama31_chatbot_rep/chatbot_flat/c1"),
 ("Llama-3.1-8B","r100 chatbot c1",100,"llama31","r100_chatbot_c1/chatbot_flat/c1","r100_chatbot_flat_c1/chatbot_flat/c1"),
 ("Llama-3.1-8B","r100 chatbot c2",100,"llama31","r100_chatbot_c2/chatbot_flat/c2","r100_chatbot_flat_c2/chatbot_flat/c2")]
g=lambda r,k,f='avg': (r[k].get(f) if isinstance(r[k],dict) else r[k]) if k in r else None

def rows_for(L,E,req):
    def n(v,d=2): return "—" if v is None else f"{v:,.{d}f}"
    def ratio(a,b,inv=False):
        if not a or not b: return "—"
        r=(b/a) if not inv else (a/b)
        return f"{r:.2f}×"
    return [
     ("TTFT  avg (ms)",           n(g(L,'time_to_first_token'),1),      n(g(E,'time_to_first_token'),1),      ratio(g(L,'time_to_first_token'),g(E,'time_to_first_token'))),
     ("ITL = TPOT  avg (ms)",     n(g(L,'inter_token_latency')),        n(g(E,'inter_token_latency')),        ratio(g(L,'inter_token_latency'),g(E,'inter_token_latency'))),
     ("E2E latency  avg (ms)",    n(g(L,'request_latency'),1),          n(g(E,'request_latency'),1),          ratio(g(L,'request_latency'),g(E,'request_latency'))),
     ("Prefill throughput (tok/s/user)", n(g(L,'prefill_throughput_per_user'),1), n(g(E,'prefill_throughput_per_user'),1), ratio(g(L,'prefill_throughput_per_user'),g(E,'prefill_throughput_per_user'),inv=True)),
     ("Decode throughput (tok/s/user)",  n(g(L,'output_token_throughput_per_user')), n(g(E,'output_token_throughput_per_user')), ratio(g(L,'output_token_throughput_per_user'),g(E,'output_token_throughput_per_user'),inv=True)),
     ("Output throughput, system (tok/s)", n(g(L,'output_token_throughput')), n(g(E,'output_token_throughput')), ratio(g(L,'output_token_throughput'),g(E,'output_token_throughput'),inv=True)),
     ("Request throughput (req/s)",n(g(L,'request_throughput'),4),      n(g(E,'request_throughput'),4),       ratio(g(L,'request_throughput'),g(E,'request_throughput'),inv=True)),
     ("ISL  avg / total (tokens)", f"{g(L,'input_sequence_length'):,.1f} / {g(L,'total_isl'):,.0f}",
                                   f"{g(E,'input_sequence_length'):,.1f} / {g(E,'total_isl'):,.0f}",
                                   "相同" if g(L,'total_isl')==g(E,'total_isl') else "不同"),
     ("OSL  avg / total (tokens)", f"{g(L,'output_sequence_length'):,.1f} / {g(L,'total_osl'):,.0f}",
                                   f"{g(E,'output_sequence_length'):,.1f} / {g(E,'total_osl'):,.0f}",
                                   "相同" if g(L,'total_osl')==g(E,'total_osl') else "不同"),
     ("Request number",            f"{req}", f"{req}", ""),
     ("Successful requests",       f"{g(L,'request_count'):,.0f} / {req}", f"{g(E,'request_count'):,.0f} / {req}", ""),
    ]

def add_table_slides(prs, blank, head, tb, rule):
    for mdl,name,req,mk,l,e in PAIRS:
        L=json.load(open(f"results/nvfp4/spark_repro_{mk}/{l}/profile_export_aiperf.json"))
        E=json.load(open(f"results/nvfp4/{e}/profile_export_aiperf.json"))
        s=blank(); head(s,f"{mdl} · {name}","附錄：逐組對照")
        data=rows_for(L,E,req)
        rowsn=len(data)+1
        shp=s.shapes.add_table(rowsn,4,Inches(0.9),Inches(2.0),Inches(11.5),Inches(0.365*rowsn))
        t=shp.table
        t.columns[0].width=Inches(4.0); t.columns[1].width=Inches(2.9)
        t.columns[2].width=Inches(2.9); t.columns[3].width=Inches(1.7)
        hdr=["指標","spark 本機 vLLM","Dynamo 端點","端點 / 本機"]
        for c,txt in enumerate(hdr):
            cell=t.cell(0,c); cell.text=""
            p=cell.text_frame.paragraphs[0]; r=p.add_run(); r.text=txt
            r.font.size=Pt(13); r.font.bold=True; r.font.name=FONT
            r.font.color.rgb=SPARK if c==1 else (ENDPT if c==2 else INK)
            p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
            cell.fill.solid(); cell.fill.fore_color.rgb=HEADBG
        for i,row in enumerate(data,start=1):
            for c,txt in enumerate(row):
                cell=t.cell(i,c); cell.text=""
                p=cell.text_frame.paragraphs[0]; r=p.add_run(); r.text=str(txt)
                r.font.size=Pt(12.5); r.font.name=FONT; r.font.color.rgb=INK
                r.font.bold = (c==3 and txt not in ("","—"))
                if c==3 and txt=="相同": r.font.color.rgb=RGBColor(0x1b,0xaf,0x7a)
                p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
                cell.fill.solid(); cell.fill.fore_color.rgb=SURF
        y=2.0+0.365*rowsn+0.2
        tb(s,0.9,y,11.5,0.8,
           "ISL / OSL 兩列為 parity 檢查：同一組數據在兩邊逐 percentile 完全相同，代表比的是同一件事。\n"
           "「端點 / 本機」欄位：延遲類為端點慢幾倍，吞吐類為本機是端點的幾倍。\n"
           "Prefill / Decode 為 per-user 速率；Output throughput, system 為伺服器整體輸出。",11,False,MUTED,space=3)
