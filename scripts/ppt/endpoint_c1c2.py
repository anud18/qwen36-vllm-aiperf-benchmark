import json
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

INK=RGBColor(0x26,0x26,0x2a); MUTED=RGBColor(0x6b,0x6b,0x70); ENDPT=RGBColor(0xeb,0x68,0x34)
SPARK=RGBColor(0x2a,0x78,0xd6); SURF=RGBColor(0xfc,0xfc,0xfb); HEADBG=RGBColor(0xf2,0xf2,0xf0)
WARN=RGBColor(0xe3,0x49,0x48); GOOD=RGBColor(0x1b,0xaf,0x7a); FONT="Noto Sans CJK TC"
g=lambda r,k,f='avg': (r[k].get(f) if isinstance(r[k],dict) else r[k]) if k in r else None

def add_c1c2_slides(prs, blank, head, tb, rule):
    A=json.load(open("results/nvfp4/r100_chatbot_flat_c1/chatbot_flat/c1/profile_export_aiperf.json"))
    B=json.load(open("results/nvfp4/r100_chatbot_flat_c2/chatbot_flat/c2/profile_export_aiperf.json"))
    s=blank(); head(s,"Dynamo 端點：併發 1 vs 併發 2","附錄：端點自身的併發比較　·　Llama-3.1-8B · chatbot_flat · 100 請求")
    def n(v,d=2): return f"{v:,.{d}f}"
    def rat(a,b,inv=False):   # always c2 / c1
        return f"{b/a:.2f}×"
    def mk(lbl,d,key,nd):
        a,b=g(A,key),g(B,key); r=b/a
        worse=(r>1.02) if d=="↓" else ((r<0.98) if d=="↑" else False)
        return (f"{d} {lbl}", n(a,nd), n(b,nd), (f"{r:.3f}×" if r<0.1 else f"{r:.2f}×"), worse)
    rows=[mk("TTFT avg (ms)","↓","time_to_first_token",1),
     mk("ITL = TPOT avg (ms)","↓","inter_token_latency",2),
     mk("E2E latency avg (ms)","↓","request_latency",1),
     mk("Prefill throughput (tok/s/user)","↑","prefill_throughput_per_user",1),
     mk("Decode throughput (tok/s/user)","↑","output_token_throughput_per_user",2),
     mk("Output throughput, system (tok/s)","↑","output_token_throughput",2),
     mk("Request throughput (req/s)","↑","request_throughput",4),
     mk("Benchmark duration (s)","↓","benchmark_duration",1),
     mk("Prefix cache 命中率 (%)","↑","overall_usage_prompt_cache_read_pct",2),
     ("ISL total (tokens)",      f"{g(A,'total_isl'):,.0f}",      f"{g(B,'total_isl'):,.0f}",      "相同", False),
     ("OSL total (tokens)",      f"{g(A,'total_osl'):,.0f}",      f"{g(B,'total_osl'):,.0f}",      "相同", False),
     ("Successful requests",     f"{g(A,'request_count'):.0f} / 100", f"{g(B,'request_count'):.0f} / 100", "", False),
    ]
    nrow=len(rows)+1
    t=s.shapes.add_table(nrow,4,Inches(0.9),Inches(1.9),Inches(11.5),Inches(0.325*nrow)).table
    t.columns[0].width=Inches(4.3); t.columns[1].width=Inches(2.5)
    t.columns[2].width=Inches(2.5); t.columns[3].width=Inches(2.2)
    for c,txt in enumerate(["指標　（↓ 越小越好　↑ 越大越好）","併發 1","併發 2","c2 ÷ c1"]):
        cell=t.cell(0,c); cell.text=""
        p=cell.text_frame.paragraphs[0]; r=p.add_run(); r.text=txt
        r.font.size=Pt(12.5); r.font.bold=True; r.font.name=FONT; r.font.color.rgb=ENDPT if c in (1,2) else INK
        p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
        cell.fill.solid(); cell.fill.fore_color.rgb=HEADBG
    for i,rw in enumerate(rows,start=1):
        worse=rw[4]
        for c,txt in enumerate(rw[:4]):
            cell=t.cell(i,c); cell.text=""
            p=cell.text_frame.paragraphs[0]; r=p.add_run(); r.text=str(txt)
            r.font.size=Pt(12.5); r.font.name=FONT; r.font.color.rgb=INK
            r.font.bold=(c==3 and txt not in ("",))
            if c==3 and txt not in ("",): r.font.color.rgb = WARN if worse else GOOD
            p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
            cell.fill.solid(); cell.fill.fore_color.rgb=SURF
    y=1.9+0.325*nrow+0.18
    tb(s,0.9,y,11.5,0.5,
       "總工作量固定 100 個請求，併發加倍後端點卻多花 13.6% 的時間才跑完 —— 加大併發是負效益。",13,True,INK)
    tb(s,0.9,y+0.34,11.5,0.4,
       "比值欄一律為 c2 ÷ c1。紅色 = 併發加倍後變差，綠色 = 持平或變好；方向由指標本身決定。",11,False,MUTED)
    tb(s,0.9,y+0.7,11.5,0.8,
       "機制在 decode：ITL 67.4 → 181.7 ms（+170%），per-user decode 掉到 6.62 tok/s，系統輸出從 8.97 降到 7.89。\n"
       "同一組測試在 spark 本機是加速 2.09×、系統輸出 13.74 → 28.75 tok/s（見第 5 頁）。",11.5,False,MUTED,space=3)

    s=blank(); head(s,"端點其餘 c1 / c2 組合為何不可比","附錄：端點自身的併發比較")
    tb(s,0.9,2.1,11.5,0.5,"端點側只有上一頁那一組是乾淨的併發對照。其餘三組各有不同問題：",15,False,INK)
    items=[("Qwen3-VL · w2 chatbot c1 / c2",
            "兩點的請求數與 trace 切片不同（10 筆 entries[2:12] 對 20 筆 entries[2:22]），ISL 總和 5,937 對 12,040。\n"
            "這是兩個不同的 workload，不是同一工作量的併發比較。"),
           ("Qwen3-VL · w0 agent c1 / c2",
            "c2 的 20 筆中有 6 筆撞上 ngrok 月流量配額（403 ERR_NGROK_725），實際只完成 14 筆，\n"
            "ISL 總和因此變成 21,714 對 c1 的 33,160。工作量不同，無法比較。"),
           ("Llama-3.1-8B · r100 agent c1 / c2",
            "兩點分別遺失 3 筆與 10 筆 Cloudflare 524 逾時，ISL 總和 155,303 對 135,900。\n"
            "524 砍掉的是最慢的請求，兩邊被砍掉的比例還不同，比較會雙重失真。")]
    yy=2.75
    for title,body in items:
        tb(s,0.9,yy,11.5,0.4,title,14,True,INK)
        tb(s,0.9,yy+0.42,11.5,0.8,body,12,False,MUTED,space=2)
        yy+=1.32
    rule(s,0.9,yy+0.05,11.5)
    tb(s,0.9,yy+0.3,11.5,0.6,
       "spark 本機四組 c1 / c2 全部有效（ISL 逐點相同、0 錯誤），加速倍率 1.50× 至 2.09×。",13,True,INK)
