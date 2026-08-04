from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

INK=RGBColor(0x26,0x26,0x2a); MUTED=RGBColor(0x6b,0x6b,0x70); SURF=RGBColor(0xfc,0xfc,0xfb)
HEADBG=RGBColor(0xf2,0xf2,0xf0); GOOD=RGBColor(0x1b,0xaf,0x7a); WARN=RGBColor(0xe3,0x49,0x48)
SPARK=RGBColor(0x2a,0x78,0xd6); FONT="Noto Sans CJK TC"

LAT=[("TTFT","time_to_first_token","第一個 token 送達的時間 − 請求送出的時間","使用者等多久才看到第一個字","↓"),
 ("ITL = TPOT","inter_token_latency","(E2E 延遲 − TTFT) ÷ (OSL − 1)","出字期間平均每個 token 的間隔","↓"),
 ("E2E latency","request_latency","請求結束時間 − 請求送出時間","一個請求從頭到尾的總時間","↓"),
 ("Prefill throughput","prefill_throughput_per_user","ISL ÷ TTFT","讀取 prompt 的速度（每個使用者）","↑"),
 ("Decode throughput","output_token_throughput_per_user","1000 ÷ ITL","出字速度（每個使用者感受到的）","↑")]
AGG=[("Output throughput, system","output_token_throughput","OSL 總和 ÷ 測試總時長","伺服器整體每秒產出多少 token","↑"),
 ("Request throughput","request_throughput","請求數 ÷ 測試總時長","伺服器每秒消化多少個請求","↑"),
 ("Benchmark duration","benchmark_duration","最後一筆結束 − 最早一筆送出","整輪測試的實際牆鐘時間","↓"),
 ("ISL","input_sequence_length","伺服器回報的 usage.prompt_tokens","輸入長度；本次兩邊刻意設為相同","—"),
 ("OSL","output_sequence_length","伺服器回報的 usage.completion_tokens","輸出長度；由 trace 以 ignore_eos 鎖定","—"),
 ("Prefix cache 命中率","overall_usage_prompt_cache_read_pct","快取命中的 prompt token ÷ 全部 prompt token","有多少輸入是重用快取、不必重算","↑"),
 ("Successful requests","request_count","實際完成且被統計的請求數","未完成者多為 tunnel 層失敗，非伺服器問題","↑")]

def _table(s, rows, top, tb):
    n=len(rows)+1
    t=s.shapes.add_table(n,4,Inches(0.55),Inches(top),Inches(12.2),Inches(0.335*n)).table
    t.columns[0].width=Inches(2.5); t.columns[1].width=Inches(3.5)
    t.columns[2].width=Inches(5.3); t.columns[3].width=Inches(0.9)
    for c,txt in enumerate(["指標","算法","意思","方向"]):
        cell=t.cell(0,c); cell.text=""
        p=cell.text_frame.paragraphs[0]; r=p.add_run(); r.text=txt
        r.font.size=Pt(12); r.font.bold=True; r.font.name=FONT; r.font.color.rgb=INK
        p.alignment=PP_ALIGN.CENTER if c==3 else PP_ALIGN.LEFT
        cell.fill.solid(); cell.fill.fore_color.rgb=HEADBG
    for i,(name,field,formula,meaning,d) in enumerate(rows,start=1):
        for c,txt in enumerate([name,formula,meaning,d]):
            cell=t.cell(i,c); cell.text=""
            p=cell.text_frame.paragraphs[0]; r=p.add_run(); r.text=txt
            r.font.size=Pt(11.5); r.font.name=FONT
            r.font.bold=(c==0)
            r.font.color.rgb = (GOOD if d=="↑" else WARN if d=="↓" else MUTED) if c==3 else INK
            p.alignment=PP_ALIGN.CENTER if c==3 else PP_ALIGN.LEFT
            cell.fill.solid(); cell.fill.fore_color.rgb=SURF
        cell=t.cell(i,1)
        cell.text_frame.paragraphs[0].runs[0].font.name="DejaVu Sans Mono"
    return top+0.335*n

def add_glossary_slides(prs, blank, head, tb, rule):
    s=blank(); head(s,"指標的算法與意思（一）　延遲與單使用者速率","附錄：名詞解釋")
    y=_table(s,LAT,2.0,tb)
    tb(s,0.55,y+0.25,12.2,1.4,
      "全部取自 aiperf 的欄位，未經二次加工。公式已用逐筆原始記錄驗算，以某一筆為例：\n"
      "ITL = (18,200.474 − 79.131) ÷ (252 − 1) = 72.196584 ms　·　Prefill = 69 ÷ 0.07913 = 871.97 tok/s　·　Decode = 1000 ÷ 72.196584 = 13.85 tok/s\n"
      "Prefill / Decode 皆為 per-user：描述單一使用者的體感，不是伺服器總吞吐。",12,False,MUTED,space=4)

    s=blank(); head(s,"指標的算法與意思（二）　整體吞吐與資料量","附錄：名詞解釋")
    y=_table(s,AGG,2.0,tb)
    tb(s,0.55,y+0.25,12.2,1.2,
      "per-user 與 system 是兩件事：併發加倍時，單一使用者可能變慢，但伺服器總產出可能上升 —— 兩者要一起看。\n"
      "Prefill throughput 會被 prefix cache 灌水：命中的 token 計入 ISL 卻未實際運算，故端點真實 prefill 比表上更慢。",12,False,MUTED,space=4)

    s=blank(); head(s,"表格怎麼讀：比值方向","附錄：名詞解釋")
    tb(s,0.9,2.15,11.5,0.6,"問題：有些指標越大越好，有些越小越好，比值就容易讀反。",16,True,INK)
    tb(s,0.9,2.85,11.5,1.6,
      "本簡報的作法是「除法只有一種、方向交給指標本身」：\n\n"
      "1.  比值一律是「後者 ÷ 前者」 —— 對照表為端點 ÷ 本機，併發頁為 c2 ÷ c1。不因指標而改變除法方向。\n"
      "2.  每個指標名稱前標示 ↓（越小越好）或 ↑（越大越好）。\n"
      "3.  比值用顏色標示結果：紅色 = 較差，綠色 = 持平或較好。",14,False,INK,space=8)
    rule(s,0.9,4.75,11.5)
    tb(s,0.9,5.0,11.5,0.5,"例：同一頁上這兩列都代表「端點較差」，但數字一大一小 ——",14,True,INK)
    ex=[("↓ TTFT avg (ms)","109.9","8,693.7","79.12×",True),
        ("↑ Decode throughput (tok/s/user)","30.48","20.22","0.66×",True),
        ("—  ISL total (tokens)","11,698","11,698","1.00×",False)]
    t=s.shapes.add_table(4,4,Inches(0.9),Inches(5.5),Inches(11.5),Inches(0.35*4)).table
    t.columns[0].width=Inches(4.6); t.columns[1].width=Inches(2.5)
    t.columns[2].width=Inches(2.5); t.columns[3].width=Inches(1.9)
    for c,txt in enumerate(["指標","spark 本機","Dynamo 端點","端點 ÷ 本機"]):
        cell=t.cell(0,c); cell.text=""
        p=cell.text_frame.paragraphs[0]; r=p.add_run(); r.text=txt
        r.font.size=Pt(12); r.font.bold=True; r.font.name=FONT; r.font.color.rgb=INK
        p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
        cell.fill.solid(); cell.fill.fore_color.rgb=HEADBG
    for i,(lbl,a,b,rt,worse) in enumerate(ex,start=1):
        for c,txt in enumerate([lbl,a,b,rt]):
            cell=t.cell(i,c); cell.text=""
            p=cell.text_frame.paragraphs[0]; r=p.add_run(); r.text=txt
            r.font.size=Pt(12); r.font.name=FONT; r.font.color.rgb=INK
            if c==3: r.font.bold=True; r.font.color.rgb = WARN if worse else GOOD
            p.alignment=PP_ALIGN.LEFT if c==0 else PP_ALIGN.RIGHT
            cell.fill.solid(); cell.fill.fore_color.rgb=SURF
    tb(s,0.9,6.95,11.5,0.4,
      "79.12× 是「慢了 79 倍」，0.66× 是「只剩六成六」—— 兩者都紅色，不必回頭想除法方向。",12,False,MUTED)
