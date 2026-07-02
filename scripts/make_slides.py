#!/usr/bin/env python3
"""Generate slides/benchmark_overview.pptx — benchmarking system architecture & workloads.

All numbers come from committed docs/data: DATASETS.md (scripts/dataset_stats.py),
RUNBOOK.md, FORMAT.md, report_v6/SUMMARY.md.  Re-run after changing any of those:

    pip install python-pptx && python3 scripts/make_slides.py
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ROOT = os.path.join(os.path.dirname(__file__), "..")
OUT = os.path.join(ROOT, "slides", "benchmark_overview.pptx")

# palette
INK = RGBColor(0x20, 0x28, 0x30)      # near-black text
MUT = RGBColor(0x5A, 0x66, 0x70)      # muted gray text
ACC = RGBColor(0x76, 0xB9, 0x00)      # nvidia green accent
BOX = RGBColor(0xF2, 0xF5, 0xF7)      # light box fill
HDR = RGBColor(0x20, 0x28, 0x30)      # table header fill
LNE = RGBColor(0xC9, 0xD2, 0xD8)      # box border
WHT = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def textbox(s, x, y, w, h, lines, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, leading=1.0):
    """lines: str or list of (text, opts) tuples; opts keys: size,color,bold,bullet,gap"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [(lines, {})]
    for i, (txt, o) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = o.get("align", align)
        p.line_spacing = leading
        p.space_after = Pt(o.get("gap", 2))
        r = p.add_run()
        r.text = ("•  " + txt) if o.get("bullet") else txt
        f = r.font
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
        f.name = "Calibri"
    return tb


def box(s, x, y, w, h, fill=BOX, line=LNE, dash=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.055
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(1.1)
    if dash:
        d = sh.line._get_or_add_ln()
        pd = d.makeelement(qn("a:prstDash"), {"val": dash})
        d.append(pd)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def boxtext(s, x, y, w, h, title, items, fill=BOX, tsize=13, isize=10.5, dash=None,
            tcolor=INK):
    box(s, x, y, w, h, fill=fill, dash=dash)
    lines = [(title, {"size": tsize, "bold": True, "color": tcolor, "gap": 4})]
    lines += [(t, {"bullet": True, "size": isize, "color": MUT, "gap": 2}) for t in items]
    textbox(s, x + 0.14, y + 0.08, w - 0.28, h - 0.16, lines)


def harrow(s, x, y, w, label=None, above=True, rev=False):
    shp = s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW if rev else MSO_SHAPE.RIGHT_ARROW,
                             Inches(x), Inches(y), Inches(w), Inches(0.16))
    shp.adjustments[0] = 0.6
    shp.adjustments[1] = 0.55
    shp.fill.solid(); shp.fill.fore_color.rgb = ACC
    shp.line.fill.background(); shp.shadow.inherit = False
    if label:
        ly = y - 0.26 if above else y + 0.18
        textbox(s, x - 0.35, ly, w + 0.7, 0.3, label, size=9.5, color=MUT, align=PP_ALIGN.CENTER)


def varrow(s, x, y, h, up=False):
    shp = s.shapes.add_shape(MSO_SHAPE.UP_ARROW if up else MSO_SHAPE.DOWN_ARROW,
                             Inches(x), Inches(y), Inches(0.16), Inches(h))
    shp.adjustments[0] = 0.6
    shp.adjustments[1] = 0.55
    shp.fill.solid(); shp.fill.fore_color.rgb = ACC
    shp.line.fill.background(); shp.shadow.inherit = False


def heading(s, title, sub=None):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(0.42), Inches(0.09), Inches(0.52))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background(); bar.shadow.inherit = False
    textbox(s, 0.65, 0.30, 12.2, 0.6, title, size=25, bold=True)
    if sub:
        textbox(s, 0.66, 0.88, 12.2, 0.35, sub, size=12, color=MUT)


def table(s, x, y, w, h, headers, rows, widths=None, size=10, hsize=10.5, align0=True):
    shp = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    t = shp.table
    if widths:
        total = sum(widths)
        for i, cw in enumerate(widths):
            t.columns[i].width = Emu(int(Inches(w) * cw / total))
    for j, htxt in enumerate(headers):
        c = t.cell(0, j)
        c.text = htxt
        c.fill.solid(); c.fill.fore_color.rgb = HDR
        c.margin_top = c.margin_bottom = Pt(2)
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(hsize); r.font.bold = True; r.font.color.rgb = WHT; r.font.name = "Calibri"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i + 1, j)
            c.text = str(cell)
            c.fill.solid()
            c.fill.fore_color.rgb = WHT if i % 2 == 0 else BOX
            c.margin_top = c.margin_bottom = Pt(2)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size); r.font.name = "Calibri"; r.font.color.rgb = INK
                    if j == 0 and align0:
                        r.font.bold = True
    return t


# ================================================================ 1. title
s = slide()
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
bg.fill.solid(); bg.fill.fore_color.rgb = INK; bg.line.fill.background(); bg.shadow.inherit = False
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.62), Inches(0.12), Inches(1.5))
bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background(); bar.shadow.inherit = False
textbox(s, 1.2, 2.45, 11.0, 1.2, "LLM Serving Benchmark", size=42, bold=True, color=WHT)
textbox(s, 1.2, 3.35, 11.0, 0.7, "System architecture & workloads", size=24, color=ACC)
textbox(s, 1.2, 4.45, 11.0, 1.2, [
    ("Qwen3.6-35B-A3B-FP8  ·  vLLM 0.22.1  ·  NVIDIA GB10 (DGX Spark)  ·  NVIDIA aiperf", {"size": 15, "color": WHT, "gap": 6}),
    ("5 workloads from real datasets & production traces: chatbot · coding · RAG · agent · toolagent", {"size": 13, "color": RGBColor(0xB9, 0xC4, 0xCC)}),
])

# ================================================================ 2. architecture
s = slide()
heading(s, "System architecture", "one client, one server, everything measured — scripts/run_final.sh drives it end-to-end")

boxtext(s, 0.45, 1.32, 12.45, 0.86, "Orchestration — scripts/run_final.sh", [
    "restart vLLM before each workload (clean cache)  ·  warmup 32 req (coding 8)  ·  concurrency sweep  ·  wedge watchdog: kill if token counters frozen 120 s, retry ≤3  ·  prefix-hit delta per point",
], tsize=12, isize=10)

boxtext(s, 0.45, 2.42, 4.0, 2.6, "aiperf 0.10.0 — load generator", [
    "Docker aiperf:local, --network host",
    "builds OpenAI chat payloads per workload file, streams responses",
    "multi-turn: replays conversations, history = model's own replies",
    "measures TTFT · ITL · latency · tok/s per request  (seed 42)",
], tsize=13, isize=10.5)

boxtext(s, 0.45, 5.22, 4.0, 1.75, "datasets/aiperf/final_*.jsonl", [
    "5 workloads from public datasets + production traces (DATASETS.md)",
    "built by scripts/build_final.py",
], tsize=12, isize=10)
varrow(s, 2.37, 5.02, 0.22, up=True)

boxtext(s, 8.45, 2.42, 4.45, 2.6, "vLLM 0.22.1 (nightly, Docker)", [
    "Qwen3.6-35B-A3B-FP8 — MoE 256 experts / 8 active, ~37 GB, 256 K ctx",
    "/v1/chat/completions (streaming) · reasoning parser qwen3",
    "prefix caching ON (hit rate exported via /metrics)",
    "max-model-len = max-num-batched-tokens = 248,320 · max-num-seqs 32 · GPU util 0.85 → KV ≈ 1.3 M tok",
], tsize=13, isize=10.5)

boxtext(s, 8.45, 5.22, 4.45, 0.8, "NVIDIA GB10 (DGX Spark)", [
    "Grace-Blackwell · 128 GB unified memory · aarch64 · CUDA 13",
], tsize=11.5, isize=9.5)

harrow(s, 4.6, 2.9, 3.7, label="HTTP /v1/chat/completions (streaming)")
harrow(s, 4.6, 3.62, 3.7, label="SSE token stream", above=False, rev=True)

boxtext(s, 4.6, 4.42, 3.7, 0.95, "optional tracing plane (docs/OBSERVABILITY.md)", [
    "LiteLLM SDK/gateway → traces/litellm_trace.jsonl — per-request content / tokens / cost",
], tsize=10, isize=9, dash="dash")

boxtext(s, 4.6, 5.55, 3.7, 1.42, "Prometheus :9090 → Grafana :3000", [
    "scrapes vLLM /metrics every 2 s · official vLLM dashboard",
    "run_final.sh reads the same counters → prefix.json per point",
], tsize=11.5, isize=9.5)
varrow(s, 8.05, 5.0, 0.5, up=True)

# ================================================================ 3. workloads & sources
s = slide()
heading(s, "Five workloads, five real data sources", "each targets a different serving regime — see DATASETS.md for provenance & stats")
table(
    s, 0.45, 1.5, 12.45, 4.6,
    ["Workload", "Original dataset", "v6 benchmark file", "aiperf type", "Scenario it models"],
    [
        ["chatbot", "ShareGPT V3 — 94,145 real user↔ChatGPT conversation records (Vicuna cleaned split)",
         "aiperf built-in pool (73,277 sessions / 252,196 turns)", "public",
         "open-domain multi-turn chat; medium in / medium out"],
        ["coding", "Inferact/codex_swebenchpro_traces — 610 Codex agent traces solving SWE-bench Pro (11 OSS repos)",
         "final_coding.jsonl\n100 conv × 6 turns", "multi_turn",
         "coding agent: ≈12 K-token repo preamble + tool outputs, history re-sent every turn"],
        ["rag", "yixuantt/MultiHopRAG — 2,556 multi-hop queries over 609 news articles (Sep–Dec 2023)",
         "final_rag.jsonl\n500 requests", "single_turn",
         "RAG QA: stuffed multi-doc context, few-word answer"],
        ["agent", "AI45Research/ATBench-Claw — 500 OpenClaw agent-safety trajectories (tools, skills, thinking)",
         "final_agent.jsonl\n200 conv / 887 turns", "multi_turn",
         "tool-using assistant: short mixed user + tool-result turns"],
        ["toolagent", "Mooncake FAST'25 trace (Kimi production, 1 h) — 23,608 requests; lengths + prefix-block hashes, no text",
         "final_toolagent.jsonl\n200 requests", "mooncake_trace",
         "production tool/agent traffic replay with realistic prefix reuse"],
    ],
    widths=[1.15, 4.0, 1.85, 1.55, 3.9], size=10.5, hsize=11,
)
textbox(s, 0.45, 6.45, 12.45, 0.6, [
    ("Output capped at 5,000 tokens per request (toolagent keeps the trace's own lengths; chatbot per-turn = original reply length).  "
     "Reasoning/thinking ON only for coding — the other four disable it via chat_template_kwargs.", {"size": 10.5, "color": MUT}),
])

# ================================================================ 4. measured characteristics
s = slide()
heading(s, "Workload characteristics (measured)", "file stats: scripts/dataset_stats.py (tokens ≈ chars/4) · served ISL/OSL: aiperf, report_v6")
table(
    s, 0.45, 1.5, 12.45, 3.9,
    ["Workload", "Turns / conv", "Input tokens (p50 / max)", "Served ISL avg", "Served OSL avg", "Output profile"],
    [
        ["chatbot", "multi (2–6+)", "a few hundred / turn", "366–511", "217–235", "conversational replies"],
        ["coding", "6 (fixed)", "turn-0 preamble 12.4 K / 14.1 K\ncumulative 37 K / 94 K per conv", "19.8 K – 24.2 K", "2.4 K – 4.1 K", "reasoning ON → longest decode"],
        ["rag", "1", "5.7 K / 30.1 K", "6.3 K", "≈ 3", "few-word answer (EOS ≪ cap)"],
        ["agent", "2–7 (median 4)", "turn-0 1.1 K / 8.3 K; new text / turn p50 127", "1.8 K – 3.4 K", "234–601", "short-medium tool-driven replies"],
        ["toolagent", "1 (trace replay)", "6.6 K / 120.6 K (exact, heavy-tailed)", "10.2 K", "≈180–190", "trace-defined, ≤ 929"],
    ],
    widths=[1.2, 1.5, 3.6, 1.6, 1.5, 3.0], size=10.5, hsize=11,
)
textbox(s, 0.45, 5.65, 12.45, 1.5, [
    ("Why these five — each stresses a different part of the serving stack:", {"size": 12, "bold": True, "gap": 4}),
    ("rag = pure prefill / TTFT   ·   coding = KV capacity + within-conversation prefix reuse + reasoning decode   ·   "
     "agent = many small multi-turn requests (scheduling)   ·   toolagent = block-level cross-request prefix reuse, heavy-tailed inputs   ·   "
     "chatbot = balanced baseline", {"size": 11, "color": MUT, "gap": 4}),
])

# ================================================================ 5. payload mechanics
s = slide()
heading(s, "What the model actually receives", "verified from aiperf's recorded request bodies — FORMAT.md shows real payload excerpts")
boxtext(s, 0.45, 1.45, 6.05, 2.5, "Common path", [
    "every request → vLLM /v1/chat/completions; JSONL text becomes one user message (Qwen chat template)",
    "output_length → max_completion_tokens (generation cap)",
    "extra → merged into body: chat_template_kwargs.enable_thinking steers reasoning on/off per request",
], tsize=13, isize=11)
boxtext(s, 0.45, 4.15, 6.05, 2.8, "multi_turn (coding, agent)", [
    "1 conversation line → 1 request per turn, sent in order",
    "aiperf prepends history; assistant turns are the model's own previous replies — dataset gpt text is never sent",
    "coding: the 12 K-token preamble rides along every turn ⇒ prefix-cache gold (hit 45–68 %)",
], tsize=13, isize=11)
boxtext(s, 6.85, 1.45, 6.05, 2.5, "single_turn (rag)", [
    "1 line = 1 request: instruction + full evidence articles + question",
    "context 2.8 K–30 K tokens, answers ~3 tokens ⇒ prefill-dominated",
], tsize=13, isize=11)
boxtext(s, 6.85, 4.15, 6.05, 2.8, "mooncake_trace (toolagent)", [
    "trace has no text — aiperf synthesizes prompts to the exact input_length",
    "each hash_id deterministically maps to a byte-identical 512-token block ⇒ replays production prefix-reuse (26.9 % block re-references)",
    "timestamps stripped in v6 so --concurrency applies (kept ⇒ fixed-schedule replay)",
], tsize=13, isize=11)

# ================================================================ 6. methodology
s = slide()
heading(s, "Run methodology (v6)", "designed for steady-state, cache-honest, resumable measurements")
boxtext(s, 0.45, 1.45, 6.05, 2.6, "Isolation & warmup", [
    "fresh vLLM server before each workload — no cross-workload cache leak or backlog (restart ≈ 11 min, AOT compile cache mounted)",
    "warmup 32 requests per point (coding 8: reasoning-ON warmup would eat the budget) → measured window is steady-state",
], tsize=13, isize=11)
boxtext(s, 0.45, 4.25, 6.05, 2.6, "Sweep & termination", [
    "concurrency 4 / 8 / 16 (+32 for chatbot, rag) — 17 points total",
    "chatbot · rag · toolagent → 160 requests per point",
    "agent · coding → time-boxed: 300 s + 300 s grace (long multi-turn sessions)",
], tsize=13, isize=11)
boxtext(s, 6.85, 1.45, 6.05, 2.6, "Cache accounting", [
    "prefix-cache hit rate recorded per (workload, concurrency): delta of vllm:prefix_cache_hits/queries_total → prefix.json",
    "server restart between workloads keeps these numbers attributable",
], tsize=13, isize=11)
boxtext(s, 6.85, 4.25, 6.05, 2.6, "Reliability & outputs", [
    "watchdog polls /metrics every 15 s; kills a run wedged 120 s (vLLM EngineCore deadlock), restarts server, retries ≤3 — struck 3×, all recovered",
    "results/final/<wl>/c<n>/ → summarize_final.py (metrics table incl. prefix-hit) + plot_pareto.py (figures)",
], tsize=13, isize=11)

# ================================================================ 7. key results
s = slide()
heading(s, "What the v6 run shows", "full tables: report_v6/SUMMARY.md · figures: report_v6/figures/")
textbox(s, 0.45, 1.42, 6.5, 5.6, [
    ("TTFT rises with concurrency everywhere — worst where prefill dominates:", {"size": 12.5, "bold": True, "gap": 1}),
    ("rag 3.6 s → 28.6 s (c4→c32, req/s flat ≈ 0.9: pure prefill queueing); chatbot only 0.25 s → 0.51 s", {"size": 11.5, "color": MUT, "gap": 7}),
    ("System vs per-user throughput trade:", {"size": 12.5, "bold": True, "gap": 1}),
    ("chatbot 109 → 239 out-tok/s while per-user falls 28 → 8.7 tok/s (c4→c32)", {"size": 11.5, "color": MUT, "gap": 7}),
    ("Prefix-cache hit tracks context sharing:", {"size": 12.5, "bold": True, "gap": 1}),
    ("coding 45→68 %, agent 53–59 %, toolagent ≈ 21 %, rag 9–14 %, chatbot decays 17→3 %", {"size": 11.5, "color": MUT, "gap": 7}),
    ("Prefix caching pays: re-sending coding's ~26 K-token history cost ~3 s TTFT instead of ~25 s cold (pass-2 measurement, thinking off)", {"size": 12.5, "bold": True, "gap": 7}),
    ("Big batch budget pays: max-num-batched-tokens 248,320 lets 16 rag prefills batch in one step — TTFT −35 %, req/s +24 % (vs default budget; GPU util 0.5→0.85 changed with it)", {"size": 12.5, "bold": True, "gap": 7}),
], leading=1.05)
pareto = os.path.join(ROOT, "report_v6", "figures", "pareto.png")
if os.path.exists(pareto):
    s.shapes.add_picture(pareto, Inches(7.15), Inches(1.55), width=Inches(5.9))
    textbox(s, 7.15, 6.55, 5.9, 0.4, "Throughput vs per-user interactivity (Pareto), per workload & concurrency",
            size=10, color=MUT, align=PP_ALIGN.CENTER)

# ================================================================ 8. pointers
s = slide()
heading(s, "Where to dig deeper")
table(
    s, 0.45, 1.5, 12.45, 3.4,
    ["Question", "Where"],
    [
        ["Where does each dataset come from & what does it look like?", "DATASETS.md (sources, licenses, measured stats — scripts/dataset_stats.py)"],
        ["What bytes are actually sent to the model?", "FORMAT.md / FORMAT.zh-TW.md (real recorded payloads, layer by layer)"],
        ["How do I reproduce the run?", "RUNBOOK.md (server flags, monitoring, one-command reproduce)"],
        ["Full numbers & figures?", "report_v6/ (SUMMARY.md, RESULTS.md, figures/) · earlier passes in RESULTS.md"],
        ["Per-request tracing / content capture / cost?", "docs/OBSERVABILITY.md (LiteLLM trace + cost logging)"],
    ],
    widths=[5.0, 7.45], size=11.5, hsize=12,
)
textbox(s, 0.45, 5.3, 12.45, 1.4, [
    ("Reproduce:", {"size": 12, "bold": True, "gap": 3}),
    ("python3 scripts/build_final.py   →   bash scripts/run_final.sh   →   python3 scripts/summarize_final.py && python3 scripts/plot_pareto.py",
     {"size": 11.5, "color": MUT}),
])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("wrote", os.path.abspath(OUT))
